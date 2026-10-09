"""
Renders PPTX slides to 1920x1080 PNG images using Pillow.
Reconstructs shapes, tables, textboxes, word wrapping, and pictures.
"""

import os
import io
import re
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN

FONT_REGULAR = '/System/Library/Fonts/Supplemental/Arial.ttf'
FONT_BOLD = '/System/Library/Fonts/Supplemental/Arial Bold.ttf'

SCALE = 144.0  # 13.333 inches * 144 = 1920 px, 7.5 inches * 144 = 1080 px

def pt_to_px(pt):
    return int(pt * (SCALE / 72.0))

def in_to_px(val_in):
    return int(val_in * SCALE)

def get_font(size_pt, bold=False):
    px = pt_to_px(size_pt)
    font_path = FONT_BOLD if bold else FONT_REGULAR
    try:
        return ImageFont.truetype(font_path, px)
    except Exception:
        return ImageFont.load_default()

def wrap_text(text, font, max_width):
    """Wrap text to fit inside max_width (px). Preserves explicit newlines."""
    lines = []
    paragraphs = text.split('\n')
    for p in paragraphs:
        if not p.strip():
            lines.append('')
            continue
        words = p.split(' ')
        current_line = []
        for word in words:
            test_line = ' '.join(current_line + [word])
            try:
                length = font.getlength(test_line)
            except AttributeError:
                length = font.getbbox(test_line)[2]
            if length <= max_width:
                current_line.append(word)
            else:
                if current_line:
                    lines.append(' '.join(current_line))
                    current_line = []
                try:
                    w_len = font.getlength(word)
                except AttributeError:
                    w_len = font.getbbox(word)[2]
                if w_len > max_width:
                    chunk = ""
                    for ch in word:
                        try:
                            c_len = font.getlength(chunk + ch)
                        except AttributeError:
                            c_len = font.getbbox(chunk + ch)[2]
                        if c_len <= max_width:
                            chunk += ch
                        else:
                            if chunk:
                                lines.append(chunk)
                            chunk = ch
                    if chunk:
                        current_line = [chunk]
                else:
                    current_line = [word]
        if current_line:
            lines.append(' '.join(current_line))
    return lines

def render_pptx(pptx_path, out_dir):
    os.makedirs(out_dir, exist_ok=True)
    prs = Presentation(pptx_path)
    
    slide_w_px = int(prs.slide_width.inches * SCALE)
    slide_h_px = int(prs.slide_height.inches * SCALE)
    print(f"Canvas size: {slide_w_px} x {slide_h_px}")
    
    for s_idx, slide in enumerate(prs.slides, 1):
        img = Image.new('RGB', (slide_w_px, slide_h_px), color=(18, 22, 26))
        draw = ImageDraw.Draw(img)
        
        for shape in slide.shapes:
            left = in_to_px(shape.left.inches if shape.left else 0)
            top = in_to_px(shape.top.inches if shape.top else 0)
            w = in_to_px(shape.width.inches if shape.width else 0)
            h = in_to_px(shape.height.inches if shape.height else 0)
            right = left + w
            bottom = top + h
            
            # 1. Background / Rectangles / Rounded Rectangles
            if shape.shape_type in [MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM]:
                fill_color = None
                if shape.fill and shape.fill.type is not None:
                    try:
                        c = shape.fill.fore_color.rgb
                        fill_color = (c[0], c[1], c[2])
                    except Exception:
                        pass
                
                line_color = None
                if shape.line and shape.line.color and shape.line.color.type is not None:
                    try:
                        c = shape.line.color.rgb
                        line_color = (c[0], c[1], c[2])
                    except Exception:
                        pass
                
                if fill_color:
                    # Draw rectangle or rounded rect
                    if 'Rounded' in shape.name:
                        radius = int(12 * (SCALE / 96.0))
                        draw.rounded_rectangle([left, top, right, bottom], radius=radius, fill=fill_color, outline=line_color, width=2 if line_color else 0)
                    else:
                        draw.rectangle([left, top, right, bottom], fill=fill_color, outline=line_color, width=1 if line_color else 0)

            # 2. Pictures
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    img_bytes = shape.image.blob
                    pic = Image.open(io.BytesIO(img_bytes)).convert('RGBA')
                    pic = pic.resize((w, h), Image.Resampling.LANCZOS)
                    img.paste(pic, (left, top), pic)
                    # Border around picture
                    draw.rectangle([left, top, right, bottom], outline=(40, 50, 61), width=1)
                except Exception as e:
                    print(f"  Error rendering picture: {e}")

            # 3. Tables
            elif shape.has_table:
                table = shape.table
                cur_y = top
                for r_idx, row in enumerate(table.rows):
                    row_h = in_to_px(row.height.inches if row.height else 0.4)
                    cur_x = left
                    for c_idx, cell in enumerate(row.cells):
                        col_w = in_to_px(table.columns[c_idx].width.inches if table.columns[c_idx].width else 1.0)
                        
                        # Cell background
                        cell_fill = (18, 22, 26)
                        if cell.fill and cell.fill.type is not None:
                            try:
                                c = cell.fill.fore_color.rgb
                                cell_fill = (c[0], c[1], c[2])
                            except Exception:
                                pass
                        
                        draw.rectangle([cur_x, cur_y, cur_x + col_w, cur_y + row_h], fill=cell_fill, outline=(40, 50, 61), width=1)
                        
                        # Cell text
                        text = cell.text.strip()
                        if text:
                            is_header = (r_idx == 0)
                            font = get_font(11 if is_header else 10, bold=is_header)
                            color = (0, 210, 196) if is_header else ((255, 255, 255) if c_idx == 0 else (160, 174, 192))
                            lines = wrap_text(text, font, col_w - 20)
                            line_y = cur_y + 8
                            for line in lines:
                                draw.text((cur_x + 10, line_y), line, font=font, fill=color)
                                line_y += pt_to_px(font.size) + 4
                                
                        cur_x += col_w
                    cur_y += row_h

            # 4. Text Boxes / Frames
            if shape.has_text_frame:
                tf = shape.text_frame
                margin_l = in_to_px(tf.margin_left.inches if tf.margin_left else 0.05)
                margin_t = in_to_px(tf.margin_top.inches if tf.margin_top else 0.05)
                margin_r = in_to_px(tf.margin_right.inches if tf.margin_right else 0.05)
                max_txt_w = w - (margin_l + margin_r)
                
                cur_ty = top + margin_t
                for p in tf.paragraphs:
                    txt = p.text.replace('\x0b', ' ').replace('\r', '')
                    if not txt.strip():
                        cur_ty += int(10 * (SCALE / 96.0))
                        continue
                    
                    # Font attributes
                    f_size = p.font.size.pt if (p.font and p.font.size) else 12
                    f_bold = p.font.bold if (p.font and p.font.bold is not None) else False
                    f_color = (255, 255, 255)
                    if p.font and p.font.color and p.font.color.type is not None:
                        try:
                            c = p.font.color.rgb
                            f_color = (c[0], c[1], c[2])
                        except Exception:
                            pass
                    elif p.runs:
                        r0 = p.runs[0]
                        if r0.font and r0.font.size:
                            f_size = r0.font.size.pt
                        if r0.font and r0.font.bold is not None:
                            f_bold = r0.font.bold
                        if r0.font and r0.font.color and r0.font.color.type is not None:
                            try:
                                c = r0.font.color.rgb
                                f_color = (c[0], c[1], c[2])
                            except Exception:
                                pass
                    
                    font = get_font(f_size, bold=f_bold)
                    lines = wrap_text(txt, font, max_txt_w)
                    line_height = pt_to_px(f_size) + int(4 * (SCALE / 96.0))
                    
                    for line in lines:
                        try:
                            line_len = font.getlength(line)
                        except AttributeError:
                            line_len = font.getbbox(line)[2]
                            
                        if p.alignment == PP_ALIGN.CENTER:
                            line_x = left + margin_l + (max_txt_w - line_len) // 2
                        elif p.alignment == PP_ALIGN.RIGHT:
                            line_x = left + margin_l + (max_txt_w - line_len)
                        else:
                            line_x = left + margin_l
                            
                        draw.text((line_x, cur_ty), line, font=font, fill=f_color)
                        cur_ty += line_height
                        
                    if p.space_after:
                        cur_ty += pt_to_px(p.space_after.pt)
                    if p.space_before:
                        cur_ty += pt_to_px(p.space_before.pt)

        out_file = os.path.join(out_dir, f"slide_{s_idx:02d}.png")
        img.save(out_file, 'PNG')
        print(f"Rendered slide {s_idx:02d} -> {out_file}")

if __name__ == '__main__':
    render_pptx('Final Project/Final_Presentation.pptx', 'scratch_slides')
